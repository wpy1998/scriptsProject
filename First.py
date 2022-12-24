import docx
import csv

document = docx.Document()
table = document.add_table(163, 11)
tag = 0
csv_reader = csv.reader(open("C:\\Users\\12514\\Desktop\\test.csv"))
for line in csv_reader:
    print(tag)
    if tag > 161:
        break
    table.cell(tag, 0).text = line[0]
    table.cell(tag, 1).text = line[1]
    table.cell(tag, 2).text = '郑州京拓'
    table.cell(tag, 3).text = line[2]
    table.cell(tag, 6).text = line[4]
    table.cell(tag, 7).text = line[9]
    table.cell(tag, 8).text = line[10]
    table.cell(tag, 9).text = '合格'
    tag = tag + 1

document.save('result.docx')