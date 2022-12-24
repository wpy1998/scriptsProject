import docx
import csv

# document = docx.Document()
# table = document.add_table(163, 11)
# tag = 0
# csv_reader = csv.reader(open("C:\\Users\\12514\\Desktop\\test.csv"))
# for line in csv_reader:
#     print(tag)
#     if tag > 161:
#         break
#     table.cell(tag, 0).text = line[0]
#     table.cell(tag, 1).text = line[1]
#     table.cell(tag, 2).text = '郑州京拓'
#     table.cell(tag, 3).text = line[2]
#     table.cell(tag, 6).text = line[4]
#     table.cell(tag, 7).text = line[9]
#     table.cell(tag, 8).text = line[10]
#     table.cell(tag, 9).text = '合格'
#     tag = tag + 1
#
# document.save('result.docx')

csv_reader = csv.reader(open("1.csv", encoding='utf-8'))
tag = -1
data = []
for line in csv_reader:
    son_data = []
    for i in range(13):
        son_data.append(line[i])
    tag = tag + 1
    if(tag != 0 and len(line[12]) != 0):
        result = ''
        for i in range(len(line[12])):
            if(line[12][i] == '-' and len(result) != 0):
                son_data.append(result)
                result = ''
            elif(line[12][i] >= '0' and line[12][i] <= '9'):
                result = result + line[12][i]
            elif(line[12][i] == ' '):
                continue
            else:
                result = ''

        # print(son_data)
    data.append(son_data)

csv_writer = csv.writer(open('1_1.csv', 'w', encoding='utf-8', newline=''))
max = 0
for line in range(len(data)):
    if(len(data[line]) > max):
        max = len(data[line])
    print(data[line])
csv_writer.writerows(data)
print(max)